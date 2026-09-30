"""Export every demo login to Excel.

    python -m scripts.export_accounts            # -> data/demo_accounts.xlsx

Passwords are deterministic (see app/accounts.demo_password), so this sheet
matches the database created on any machine for the default scenario (seed 42).
DEMO ONLY: fictional institutions and people; never reuse these passwords.
"""
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.accounts import GOVT_USERNAME, demo_password, get_accounts
from app.config import REPO_DIR
from app.service import get_pipeline

NAVY, BLUE, TINT = "14202B", "1F4E8C", "E6EEF8"
RISK_FILL = {"High": "F8E1DE", "Moderate": "FBE6D6", "Lower": "F8EFCF"}
thin = Side(style="thin", color="DEDCD5")


def _sheet(wb, title, headers, rows, widths, note=None):
    ws = wb.create_sheet(title)
    r0 = 1
    if note:
        ws.cell(1, 1, note).font = Font(italic=True, color="5A6570")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers))
        r0 = 3
    for j, h in enumerate(headers, 1):
        c = ws.cell(r0, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(vertical="center")
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v)
            c.border = Border(bottom=thin)
            if headers[j - 1] == "Password":
                c.font = Font(name="Consolas", bold=True, color=BLUE)
            if headers[j - 1].startswith("Risk") and v in RISK_FILL:
                c.fill = PatternFill("solid", fgColor=RISK_FILL[v])
    for j, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + max(len(rows), 1)}"
    ws.row_dimensions[r0].height = 22
    return ws


def main(out: str | None = None):
    p = get_pipeline()
    if p.state is None:
        p.run()
    acc = get_accounts()
    acc.seed_drivers(p.events())
    ring_label = {"high": "High", "moderate": "Moderate", "low": "Lower"}
    exposure = {}
    for e in p.events():
        for a in e["peak"]["impact"].get("assets", []):
            prev = exposure.get(a["id"])
            if prev is None or "hml".index(a["ring"][0]) < "hml".index(prev[0][0]):
                exposure[a["id"]] = (a["ring"], e["id"], e["place"])
    orders = {o["unit_id"]: o for o in p.store.orders()}
    toward = {v["id"]: (v, e) for e in p.events() for v in e.get("vehicles", {}).get("toward", [])}

    wb = Workbook()
    ws = wb.active
    ws.title = "Start here"
    ws["A1"] = "WeatherPulse AI · demo logins"
    ws["A1"].font = Font(size=18, bold=True, color=NAVY)
    lines = [
        "Open http://localhost:3000 → choose your role tab → type the Username exactly as in this sheet (not case-sensitive) and the Password.",
        "Institutions log in with their full name (e.g. 'Govt. High School, Koramangala, Bengaluru'); the sign-in box suggests names as you type.",
        "Citizens: use a sample account below or tap 'Create account' to register with your mobile number.",
        "Rows marked High / Moderate / Lower are inside a risk zone in the default demo scenario, so their portal shows an active alert.",
        "DEMO DATA ONLY: all institutions, drivers and citizens are fictional. Passwords are for this prototype; never reuse them.",
    ]
    for i, t in enumerate(lines, 3):
        ws.cell(i, 1, f"• {t}").alignment = Alignment(wrap_text=True)
    ws.column_dimensions["A"].width = 140
    ws.cell(9, 1, "Best accounts to try first").font = Font(bold=True, size=13, color=NAVY)
    quick = [("Government official", GOVT_USERNAME, demo_password(GOVT_USERNAME), "Full command dashboard + AI Copilot")]
    for kind, label in (("school", "School / college"), ("hospital", "Hospital")):
        best = sorted((a for a in p.assets if a["id"] in exposure and a["kind"] == kind),
                      key=lambda a: "hml".index(exposure[a["id"]][0][0]))
        if best:
            quick.append((label, best[0]["name"], demo_password(best[0]["name"]), f"{ring_label[exposure[best[0]['id']][0]]} risk · {exposure[best[0]['id']][2]}"))
    if orders:
        u = p.asset_by_id[sorted(orders, key=lambda k: "BLR" not in k)[0]]
        quick.append(("Rescue team", u["name"], demo_password(u["name"]), "Has a pre-position order"))
    bus = sorted((x for x in toward.values() if x[0]["kind"] == "bus"),
                 key=lambda x: (x[1]["location"]["nearest_city"] != "Bengaluru", x[0]["status"] != "approaching"))
    if bus:
        quick.append(("Bus driver", bus[0][0]["id"], demo_password(bus[0][0]["id"]), f"Heading into zone · {bus[0][1]['place']}"))
    quick.append(("Citizen (sample)", "9000000001", demo_password("citizen-9000000001"), "Ananya Rao · Koramangala, Bengaluru"))
    ws2 = ws
    hdr = ["Role", "Username", "Password", "What you will see"]
    for j, h in enumerate(hdr, 1):
        c = ws2.cell(10, j, h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
    for i, row in enumerate(quick, 11):
        for j, v in enumerate(row, 1):
            c = ws2.cell(i, j, v)
            if j == 3:
                c.font = Font(name="Consolas", bold=True, color=BLUE)
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 52
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 60

    _sheet(wb, "Government", ["Username", "Password", "Role", "Organisation"],
           [[GOVT_USERNAME, demo_password(GOVT_USERNAME), "Government official", "District Disaster Management Authority"]], [22, 18, 24, 44])

    def inst_rows(kind):
        rows = []
        for a in p.assets:
            if a["kind"] != kind:
                continue
            ex = exposure.get(a["id"])
            extra = [("Yes" if a["id"] in orders else "No")] if kind == "rescue_team" else []
            rows.append([a["name"], demo_password(a["name"]), a["city"], a.get("locality", ""), a["id"],
                         ring_label[ex[0]] if ex else "None", ex[1] if ex else "", *extra])
        rows.sort(key=lambda r: ({"High": 0, "Moderate": 1, "Lower": 2}.get(r[5], 3), r[2], r[0]))
        return rows
    base = ["Username (institution name)", "Password", "City", "Locality", "Institution ID", "Risk in demo", "Event"]
    note = "Sorted: institutions inside a risk zone in the demo scenario first."
    _sheet(wb, "Schools & Colleges", base, inst_rows("school"), [52, 16, 14, 22, 14, 13, 14], note)
    _sheet(wb, "Hospitals", base, inst_rows("hospital"), [52, 16, 14, 22, 14, 13, 14], note)
    _sheet(wb, "Rescue Teams", base + ["Has order"], inst_rows("rescue_team"), [52, 16, 14, 22, 14, 13, 14, 11], note)

    drv = []
    for r in acc.directory("traveller", limit=1000):
        vid = r["vehicle_id"]
        t = toward.get(vid)
        drv.append([vid, demo_password(vid), r["org"], "Yes" if t else "No",
                    (t[0]["status"] if t else ""), (t[0]["eta_min"] if t else ""), (t[1]["place"] if t else "")])
    drv.sort(key=lambda r: (r[3] != "Yes", r[0]))
    _sheet(wb, "Bus Drivers", ["Username (bus number)", "Password", "Operator", "Heading into zone", "Status", "ETA (min)", "Zone"],
           drv, [22, 16, 18, 17, 12, 10, 40], "Buses heading into a zone see the alert with a safer route; others see 'route clear'.")

    cit = [["9000000001", demo_password("citizen-9000000001"), "Ananya Rao", "Koramangala, Bengaluru"],
           ["9000000002", demo_password("citizen-9000000002"), "Rahul Das", "Bhubaneswar, Odisha"]]
    _sheet(wb, "Citizens", ["Username (mobile number)", "Password", "Name", "Home area"], cit, [24, 18, 18, 30],
           "Sample accounts. New citizens register themselves on the sign-in page (Citizen → Create account).")

    out = Path(out or REPO_DIR / "data" / "demo_accounts.xlsx")
    wb.save(out)
    counts = {s.title: s.max_row for s in wb.worksheets}
    print("saved", out, counts)
    return out


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else None)
